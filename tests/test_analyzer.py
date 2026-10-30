import csv
import pytest
from pathlib import Path
from openpyxl import Workbook
from shapes.sheets.src.analyzer import analyze_file, infer_type


class TestInferType:
    """Test the per-column type inference from a list of values."""

    def test_all_integers(self):
        assert infer_type(["1", "2", "3", "100"]) == ("integer", None)

    def test_integers_with_negatives(self):
        assert infer_type(["-1", "0", "42"]) == ("integer", None)

    def test_floats(self):
        assert infer_type(["1.5", "2.0", "3.14"]) == ("number", None)

    def test_mixed_int_and_float_is_number(self):
        assert infer_type(["1", "2.5", "3"]) == ("number", None)

    def test_booleans(self):
        assert infer_type(["true", "false", "True", "FALSE"]) == ("boolean", None)

    def test_one_non_integer_makes_string(self):
        assert infer_type(["1", "2", "abc", "4"])[0] == "string"

    def test_empty_values_skipped(self):
        assert infer_type(["1", "", "3", None, "5"]) == ("integer", None)

    def test_all_empty_is_string(self):
        assert infer_type(["", None, ""])[0] == "string"

    def test_iso_dates(self):
        result = infer_type(["2024-01-15", "2024-02-28", "2023-12-01"])
        assert result == ("date", "%Y-%m-%d")

    def test_us_dates(self):
        result = infer_type(["01/15/2024", "02/28/2024", "12/01/2023"])
        assert result == ("date", "%m/%d/%Y")

    def test_eu_dates(self):
        result = infer_type(["15.01.2024", "28.02.2024", "01.12.2023"])
        assert result == ("date", "%d.%m.%Y")

    def test_mixed_date_formats_is_string(self):
        result = infer_type(["2024-01-15", "01/15/2024"])
        assert result[0] == "string"

    def test_iso_datetimes(self):
        result = infer_type(["2024-01-15T10:30:00", "2024-02-28T08:00:00"])
        assert result[0] == "datetime"

    def test_urls_detected_as_string(self):
        result = infer_type(["https://example.com", "https://test.org"])
        assert result[0] == "string"


class TestAnalyzeFile:
    """Test full file analysis producing per-column stats."""

    def _write_csv(self, rows: list[list[str]], tmp_path: Path) -> Path:
        p = tmp_path / "test.csv"
        with open(p, "w", newline="") as f:
            writer = csv.writer(f)
            for row in rows:
                writer.writerow(row)
        return p

    def test_basic_csv(self, tmp_path):
        path = self._write_csv([
            ["Name", "Age", "Active"],
            ["Alice", "30", "true"],
            ["Bob", "25", "false"],
            ["Charlie", "35", "true"],
        ], tmp_path)
        result = analyze_file(path)

        assert result["row_count"] == 3
        assert len(result["columns"]) == 3

        name_col = result["columns"][0]
        assert name_col["header"] == "Name"
        assert name_col["position"] == 0
        assert name_col["type"] == "string"
        assert name_col["unique_count"] == 3
        assert name_col["is_unique"] is True

        age_col = result["columns"][1]
        assert age_col["type"] == "integer"

        active_col = result["columns"][2]
        assert active_col["type"] == "boolean"

    def test_enum_detection(self, tmp_path):
        rows = [["Status"]] + [[s] for s in ["active", "inactive"] * 20]
        path = self._write_csv(rows, tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["type"] == "string"
        assert set(col["values"]) == {"active", "inactive"}
        assert col["is_unique"] is False

    def test_no_enum_when_too_many_unique(self, tmp_path):
        rows = [["ID"]] + [[f"item-{i}"] for i in range(50)]
        path = self._write_csv(rows, tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["values"] is None

    def test_url_detection(self, tmp_path):
        path = self._write_csv([
            ["Link"],
            ["https://example.com"],
            ["https://test.org/page"],
        ], tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["looks_like_url"] is True

    def test_partial_urls_not_flagged(self, tmp_path):
        path = self._write_csv([
            ["Mixed"],
            ["https://example.com"],
            ["not a url"],
        ], tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["looks_like_url"] is False

    def test_text_vs_string(self, tmp_path):
        long_text = "x" * 250
        path = self._write_csv([
            ["Notes"],
            [long_text],
            ["short"],
        ], tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["max_length"] == 250

    def test_nulls_counted(self, tmp_path):
        path = self._write_csv([
            ["Value"],
            ["a"],
            [""],
            ["b"],
            [""],
        ], tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["nulls"] == 2
        assert col["total_count"] == 2

    def test_malformed_count(self, tmp_path):
        path = self._write_csv([
            ["Amount"],
            ["100"],
            ["200"],
            ["not-a-number"],
            ["400"],
        ], tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["type"] == "string"
        assert col["malformed_count"] == 1

    def test_date_format_recorded(self, tmp_path):
        path = self._write_csv([
            ["Date"],
            ["2024-01-15"],
            ["2024-02-28"],
        ], tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["type"] == "date"
        assert col["date_format"] == "%Y-%m-%d"

    def test_samples_capped_at_5(self, tmp_path):
        rows = [["Val"]] + [[f"v{i}"] for i in range(100)]
        path = self._write_csv(rows, tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert len(col["samples"]) == 5

    def test_primary_key_candidate(self, tmp_path):
        path = self._write_csv([
            ["ID", "Name"],
            ["1", "Alice"],
            ["2", "Bob"],
            ["3", "Charlie"],
        ], tmp_path)
        result = analyze_file(path)
        id_col = result["columns"][0]
        assert id_col["is_unique"] is True
        name_col = result["columns"][1]
        assert name_col["is_unique"] is True

    def test_file_path_recorded(self, tmp_path):
        path = self._write_csv([["A"], ["1"]], tmp_path)
        result = analyze_file(path)
        assert result["file"] == str(path)


class TestAnalyzeXlsx:
    def _write_xlsx(self, data: list[list], tmp_path: Path, sheet_name: str = "Sheet1") -> Path:
        p = tmp_path / "test.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        for row in data:
            ws.append(row)
        wb.save(p)
        return p

    def test_basic_xlsx(self, tmp_path):
        path = self._write_xlsx([
            ["Name", "Score"],
            ["Alice", 95],
            ["Bob", 87],
        ], tmp_path)
        result = analyze_file(path)
        assert result["row_count"] == 2
        assert len(result["columns"]) == 2
        assert result["columns"][0]["header"] == "Name"

    def test_xlsx_with_sheet_name(self, tmp_path):
        p = tmp_path / "multi.xlsx"
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Data"
        ws1.append(["A", "B"])
        ws1.append(["1", "2"])
        ws2 = wb.create_sheet("Other")
        ws2.append(["X"])
        ws2.append(["Y"])
        wb.save(p)

        result = analyze_file(p, sheet_name="Data")
        assert result["columns"][0]["header"] == "A"
        assert result["row_count"] == 1

    def test_xlsx_none_cells(self, tmp_path):
        path = self._write_xlsx([
            ["Val"],
            [None],
            ["hello"],
            [None],
        ], tmp_path)
        result = analyze_file(path)
        col = result["columns"][0]
        assert col["nulls"] == 2
        assert col["total_count"] == 1
