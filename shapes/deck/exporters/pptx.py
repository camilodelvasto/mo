"""Deck PPTX exporter — best-effort PowerPoint generation via python-pptx.

Generates an editable PPTX from a deck bite. Maps moui's five layouts
(cover, statement, full, split, grid) to PowerPoint slides with solid
backgrounds, positioned text, and embedded images.

HTML artifacts are captured as screenshots via headless Chromium (same
technique as the PDF exporter) and embedded as PNG images.

Visual effects (shimmer, glow, underline, stripe) are not reproduced.
PPTX is for editability, not visual fidelity.
"""
import tempfile
from pathlib import Path

import yaml

from cli.resolver import resolve_theme_dir

SLIDE_WIDTH_EMU = 12192000   # 13.333 inches
SLIDE_HEIGHT_EMU = 6858000   # 7.5 inches
MARGIN_EMU = 457200          # 0.5 inches
MARGIN_WIDE_EMU = 914400     # 1.0 inches


def _has_html_artifacts(slides: list) -> bool:
    """Check if any slide references an .html media artifact."""
    for slide in slides:
        for block in [slide.get("content", {}),
                      slide.get("left", {}).get("content", {}),
                      slide.get("right", {}).get("content", {})]:
            media = block.get("media", "")
            if isinstance(media, dict):
                media = media.get("src", "")
            if isinstance(media, str) and media.endswith(".html"):
                return True
        for cell in slide.get("cells", []):
            media = cell.get("content", {}).get("media", "")
            if isinstance(media, dict):
                media = media.get("src", "")
            if isinstance(media, str) and media.endswith(".html"):
                return True
    return False


def export(content_path: Path, shape_path: Path, corpus_root: Path, output_path: Path) -> Path:
    """Export a deck to PPTX. Returns the path to the generated file."""
    try:
        from pptx import Presentation
    except ImportError:
        raise SystemExit("PPTX export requires python-pptx.\nRun: pip install python-pptx")

    content_data = yaml.safe_load(content_path.read_text())
    slides = content_data.get("slides", [])
    theme_ref = content_data.get("theme")
    theme = _load_theme(theme_ref, content_path)

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH_EMU
    prs.slide_height = SLIDE_HEIGHT_EMU
    blank_layout = prs.slide_layouts[6]

    ctx = _ThemeContext(theme)

    # If any slide has HTML artifacts, start the server and capture them
    artifact_screenshots = {}
    if _has_html_artifacts(slides):
        artifact_screenshots = _capture_artifacts(
            content_path, shape_path, corpus_root, slides,
        )

    renderers = {
        "cover": _render_cover, "statement": _render_statement,
        "full": _render_full, "split": _render_split, "grid": _render_grid,
    }

    for idx, slide_data in enumerate(slides):
        pptx_slide = prs.slides.add_slide(blank_layout)
        layout = slide_data.get("layout", "full")
        renderers.get(layout, _render_full)(
            pptx_slide, slide_data, ctx, corpus_root, artifact_screenshots.get(idx),
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))

    # Clean up temp screenshots
    for path in artifact_screenshots.values():
        try:
            Path(path).unlink()
        except OSError:
            pass

    return output_path


def _capture_artifacts(content_path: Path, shape_path: Path, corpus_root: Path,
                       slides: list) -> dict[int, str]:
    """Capture HTML artifact slides as PNG screenshots.

    Returns a dict mapping slide index → path to temp PNG file.
    Uses the same server + Playwright technique as the PDF exporter.
    """
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        # Playwright not installed — skip artifact capture silently
        return {}

    from shapes.deck.exporters.pdf import _find_free_port, _start_server, _wait_for_server

    # Identify which slides need capture
    artifact_slides = set()
    for idx, slide in enumerate(slides):
        for block in [slide.get("content", {}),
                      slide.get("left", {}).get("content", {}),
                      slide.get("right", {}).get("content", {})]:
            media = block.get("media", "")
            if isinstance(media, dict):
                media = media.get("src", "")
            if isinstance(media, str) and media.endswith(".html"):
                artifact_slides.add(idx)
        for cell in slide.get("cells", []):
            media = cell.get("content", {}).get("media", "")
            if isinstance(media, dict):
                media = media.get("src", "")
            if isinstance(media, str) and media.endswith(".html"):
                artifact_slides.add(idx)

    if not artifact_slides:
        return {}

    port = _find_free_port()
    server = _start_server(content_path, shape_path, corpus_root, port)
    results = {}

    try:
        _wait_for_server(port)
        base_url = f"http://127.0.0.1:{port}"

        with sync_playwright() as p:
            browser = p.chromium.launch()
            page = browser.new_page(viewport={"width": 1280, "height": 720})

            # Hide UI chrome
            css = ".hamburger, .slide-counter, .sidebar, .sidebar-backdrop { display: none; }"

            for idx in sorted(artifact_slides):
                page.goto(f"{base_url}/#{idx + 1}", wait_until="load")
                page.wait_for_selector(".slide", state="visible", timeout=10000)
                page.add_style_tag(content=css)
                # Give artifact JS time to initialize and render
                page.wait_for_timeout(2000)

                tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
                tmp.close()
                page.screenshot(path=tmp.name, type="png")
                results[idx] = tmp.name

            browser.close()
    finally:
        server.shutdown()

    return results


# ---------------------------------------------------------------------------
# Theme context — resolves colors, fonts, sizes from theme tokens
# ---------------------------------------------------------------------------

def _load_theme(theme_ref: str, content_path: Path) -> dict:
    """Load resolved theme tokens from path relative to content file."""
    theme_dir = resolve_theme_dir(theme_ref, content_path)
    tokens_path = theme_dir / "tokens.yaml"
    if not tokens_path.exists():
        raise FileNotFoundError(f"tokens.yaml not found in {theme_dir}")
    return yaml.safe_load(tokens_path.read_text()) or {}


class _ThemeContext:
    """Holds resolved theme data for the export session."""

    def __init__(self, theme: dict):
        self.theme = theme
        self.palette = theme.get("palette", {})
        # Font families — resolve to installed system fonts with fallbacks
        self.font_heading = _resolve_font(theme.get("font-heading", ""))
        self.font_body = _resolve_font(theme.get("font-body", ""))
        self.weight_heading = int(theme.get("weight-heading", 700))
        self.weight_body = int(theme.get("weight-body", 400))
        self.default_statement_size = theme.get("default-statement-size", "xl")

    def resolve_color(self, value) -> str | None:
        """Resolve a color name/hex to a #RRGGBB string."""
        if not value:
            return None
        s = str(value)
        # Direct palette lookup
        if s in self.palette:
            return self.palette[s]
        # Semantic assignment lookup (e.g., color-text → dark → palette)
        assignment_key = f"color-{s}"
        assigned = self.theme.get(assignment_key)
        if assigned and assigned in self.palette:
            return self.palette[assigned]
        if assigned:
            return assigned
        # Already a hex value
        if s.startswith("#"):
            return s
        return None

    def resolve_bg_color(self, bg_val) -> str | None:
        """Resolve a background value — solid or gradient (uses 'from' color)."""
        if isinstance(bg_val, dict) and bg_val.get("from"):
            return self.resolve_color(bg_val["from"])
        return self.resolve_color(bg_val)


FONT_MAP = {
    "montserrat": "Montserrat",
    "lato": "Lato",
    "jetbrains-mono": "JetBrains Mono",
    "inter": "Inter",
    "fraunces": "Fraunces",
    "raleway": "Raleway",
    "source-serif-4": "Source Serif 4",
    "source-sans-3": "Source Sans 3",
}

# Fallbacks: theme font → system equivalent → Arial
FONT_FALLBACK = {
    "Montserrat": "Helvetica Neue",
    "Lato": "Helvetica Neue",
    "Inter": "Helvetica Neue",
    "Raleway": "Helvetica Neue",
    "Fraunces": "Georgia",
}

FALLBACK_DEFAULT = "Arial"


def _resolve_font(name: str) -> str:
    """Resolve a theme font name to the best available font for PPTX.

    Tries: exact name → known fallback → Arial.
    woff2 web fonts are not usable in PPTX, so we check if the font
    is likely available as a system font.
    """
    if not name:
        return FALLBACK_DEFAULT
    display_name = FONT_MAP.get(name.lower(), name.title())
    # Check if font is installed on this system
    if _is_font_installed(display_name):
        return display_name
    # Use a known fallback
    fallback = FONT_FALLBACK.get(display_name, FALLBACK_DEFAULT)
    return fallback


def _is_font_installed(font_name: str) -> bool:
    """Check if a font is available on the system (macOS)."""
    import subprocess
    try:
        result = subprocess.run(
            ["system_profiler", "SPFontsDataType"],
            capture_output=True, text=True, timeout=10,
        )
        return font_name.lower() in result.stdout.lower()
    except Exception:
        return False


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _hex_to_rgb(hex_color: str):
    from pptx.dml.color import RGBColor
    h = hex_color.lstrip("#")
    if len(h) == 6:
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    return None


def _set_slide_bg(pptx_slide, style: dict, ctx: _ThemeContext):
    bg_val = style.get("bg")
    if not bg_val:
        return
    color_hex = ctx.resolve_bg_color(bg_val)
    if color_hex:
        rgb = _hex_to_rgb(color_hex)
        if rgb:
            fill = pptx_slide.background.fill
            fill.solid()
            fill.fore_color.rgb = rgb


def _add_text_box(pptx_slide, text: str, left, top, width, height,
                  font_name: str = "", font_size: int = 18,
                  bold: bool = False, color_hex: str | None = None,
                  align: str | None = None, valign: str | None = None):
    """Add a text box with font family, size, color, horizontal and vertical alignment."""
    from pptx.util import Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    txBox = pptx_slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    if valign == "middle":
        tf.auto_size = None
        txBox.text_frame._txBody.bodyPr.set("anchor", "ctr")
    elif valign == "bottom":
        tf.auto_size = None
        txBox.text_frame._txBody.bodyPr.set("anchor", "b")

    p = tf.paragraphs[0]
    p.text = str(text)
    p.font.size = Pt(font_size)
    if font_name:
        p.font.name = font_name
    if bold:
        p.font.bold = True
    if color_hex:
        rgb = _hex_to_rgb(color_hex)
        if rgb:
            p.font.color.rgb = rgb
    if align == "center":
        p.alignment = PP_ALIGN.CENTER
    elif align == "right":
        p.alignment = PP_ALIGN.RIGHT


def _add_picture_preserving_ratio(pptx_slide, img_path: Path, left, top, max_width, max_height):
    """Add a picture preserving its aspect ratio within the given bounds."""
    from pptx.util import Emu
    try:
        from PIL import Image
        with Image.open(img_path) as img:
            img_w, img_h = img.size
    except Exception:
        # If Pillow can't read it, add with fixed dimensions
        try:
            pptx_slide.shapes.add_picture(str(img_path), left, top, max_width, max_height)
        except Exception:
            pass
        return

    # Calculate scale to fit within bounds
    scale_w = max_width / img_w
    scale_h = max_height / img_h
    scale = min(scale_w, scale_h)

    final_w = int(img_w * scale)
    final_h = int(img_h * scale)

    try:
        pptx_slide.shapes.add_picture(str(img_path), left, top, final_w, final_h)
    except Exception:
        pass


SIZE_MAP = {"display": 48, "xl": 44, "large": 36, "medium": 28, "small": 22}


def _add_content_block(pptx_slide, content: dict, style: dict, ctx: _ThemeContext,
                       left, top, width, corpus_root: Path):
    """Add content elements (label, heading, subheading, body, quote, attribution, media)."""
    from pptx.util import Emu

    text_color = ctx.resolve_color(style.get("color"))
    y = top

    if content.get("label"):
        label_color = ctx.resolve_color("accent") or text_color
        _add_text_box(pptx_slide, content["label"], left, y, width, Emu(350000),
                      font_name=ctx.font_body, font_size=11,
                      color_hex=label_color)
        y += Emu(350000)

    if content.get("heading"):
        hs = style.get("heading", {})
        size_name = hs.get("size", "large")
        font_size = SIZE_MAP.get(size_name, 36)
        _add_text_box(pptx_slide, content["heading"], left, y, width, Emu(900000),
                      font_name=ctx.font_heading, font_size=font_size,
                      bold=True, color_hex=text_color)
        y += Emu(900000)

    if content.get("subheading"):
        _add_text_box(pptx_slide, content["subheading"], left, y, width, Emu(450000),
                      font_name=ctx.font_body, font_size=18,
                      color_hex=text_color)
        y += Emu(450000)

    if content.get("body"):
        body = str(content["body"]).strip()
        body = body.replace("**", "").replace("*", "")
        body = body.replace("### ", "").replace("## ", "").replace("# ", "")
        # Estimate height from line count
        lines = body.count("\n") + 1
        box_h = max(Emu(600000), Emu(200000) * lines)
        box_h = min(box_h, SLIDE_HEIGHT_EMU - y - MARGIN_EMU)
        _add_text_box(pptx_slide, body, left, y, width, box_h,
                      font_name=ctx.font_body, font_size=14,
                      color_hex=text_color)
        y += box_h

    if content.get("quote"):
        _add_text_box(pptx_slide, f'\u201c{content["quote"]}\u201d', left, y, width, Emu(700000),
                      font_name=ctx.font_body, font_size=16,
                      color_hex=text_color)
        y += Emu(700000)

    if content.get("attribution"):
        _add_text_box(pptx_slide, f'\u2014 {content["attribution"]}', left, y, width, Emu(350000),
                      font_name=ctx.font_body, font_size=12,
                      color_hex=text_color)
        y += Emu(350000)

    if content.get("media"):
        src = content["media"]
        if isinstance(src, dict):
            src = src.get("src", "")
        media_path = corpus_root / src
        max_h = SLIDE_HEIGHT_EMU - y - MARGIN_EMU
        if max_h > Emu(500000):
            if media_path.exists() and media_path.suffix.lower() in (".png", ".jpg", ".jpeg", ".webp"):
                _add_picture_preserving_ratio(pptx_slide, media_path, left, y, width, max_h)
            elif media_path.suffix.lower() == ".html":
                # Artifacts: skip in PPTX — they require a browser to render
                pass

    return y


# ---------------------------------------------------------------------------
# Layout renderers
# ---------------------------------------------------------------------------

def _render_cover(pptx_slide, slide_data: dict, ctx: _ThemeContext, corpus_root: Path,
                   artifact_screenshot: str | None = None):
    from pptx.util import Emu

    style = slide_data.get("style", {})
    content = slide_data.get("content", {})
    meta = slide_data.get("meta", {})
    _set_slide_bg(pptx_slide, style, ctx)

    text_color = ctx.resolve_color(style.get("color"))
    cw = SLIDE_WIDTH_EMU - 2 * MARGIN_WIDE_EMU

    # Cover: heading + subheading centered in the middle area
    if content.get("heading"):
        size_name = style.get("heading", {}).get("size", "large")
        font_size = SIZE_MAP.get(size_name, 40)
        _add_text_box(pptx_slide, content["heading"],
                      MARGIN_WIDE_EMU, MARGIN_EMU, cw,
                      SLIDE_HEIGHT_EMU - 2 * MARGIN_EMU - Emu(600000),
                      font_name=ctx.font_heading, font_size=font_size,
                      bold=True, color_hex=text_color,
                      align="center", valign="middle")

    if content.get("subheading"):
        _add_text_box(pptx_slide, content["subheading"],
                      MARGIN_WIDE_EMU, SLIDE_HEIGHT_EMU // 2 + Emu(400000),
                      cw, Emu(600000),
                      font_name=ctx.font_body, font_size=20,
                      color_hex=text_color, align="center")

    if meta.get("author") or meta.get("date"):
        mt = " \u00b7 ".join(filter(None, [meta.get("author"), str(meta.get("date", ""))]))
        _add_text_box(pptx_slide, mt, MARGIN_WIDE_EMU,
                      SLIDE_HEIGHT_EMU - Emu(600000), cw, Emu(350000),
                      font_name=ctx.font_body, font_size=11,
                      color_hex=text_color, align="center")


def _render_statement(pptx_slide, slide_data: dict, ctx: _ThemeContext, corpus_root: Path,
                      artifact_screenshot: str | None = None):
    from pptx.util import Emu

    style = slide_data.get("style", {})
    content = slide_data.get("content", {})
    _set_slide_bg(pptx_slide, style, ctx)

    text_color = ctx.resolve_color(style.get("color"))
    size_name = style.get("heading", {}).get("size", ctx.default_statement_size)
    font_size = SIZE_MAP.get(size_name, 40)

    # Statement slides: heading centered in the full slide area
    if content.get("heading"):
        _add_text_box(pptx_slide, content["heading"],
                      MARGIN_WIDE_EMU, MARGIN_EMU,
                      SLIDE_WIDTH_EMU - 2 * MARGIN_WIDE_EMU,
                      SLIDE_HEIGHT_EMU - 2 * MARGIN_EMU,
                      font_name=ctx.font_heading, font_size=font_size,
                      bold=True, color_hex=text_color,
                      align="center", valign="middle")

    if content.get("body"):
        body = str(content["body"]).strip()
        _add_text_box(pptx_slide, body,
                      MARGIN_WIDE_EMU, SLIDE_HEIGHT_EMU * 3 // 4 - Emu(200000),
                      SLIDE_WIDTH_EMU - 2 * MARGIN_WIDE_EMU, Emu(1200000),
                      font_name=ctx.font_body, font_size=16,
                      color_hex=text_color, align="center")


def _render_full(pptx_slide, slide_data: dict, ctx: _ThemeContext, corpus_root: Path,
                 artifact_screenshot: str | None = None):
    style = slide_data.get("style", {})
    content = slide_data.get("content", {})
    _set_slide_bg(pptx_slide, style, ctx)

    # If this slide has an artifact screenshot, embed it as a full-slide image
    if artifact_screenshot:
        _add_picture_preserving_ratio(
            pptx_slide, Path(artifact_screenshot), 0, 0,
            SLIDE_WIDTH_EMU, SLIDE_HEIGHT_EMU,
        )
        return

    _add_content_block(pptx_slide, content, style, ctx,
                       MARGIN_EMU, MARGIN_EMU,
                       SLIDE_WIDTH_EMU - 2 * MARGIN_EMU, corpus_root)


def _render_split(pptx_slide, slide_data: dict, ctx: _ThemeContext, corpus_root: Path,
                  artifact_screenshot: str | None = None):
    from pptx.util import Emu

    style = slide_data.get("style", {})
    _set_slide_bg(pptx_slide, style, ctx)

    # If this slide has an artifact screenshot, embed it as full-slide
    if artifact_screenshot:
        _add_picture_preserving_ratio(
            pptx_slide, Path(artifact_screenshot), 0, 0,
            SLIDE_WIDTH_EMU, SLIDE_HEIGHT_EMU,
        )
        return

    text_color = ctx.resolve_color(style.get("color"))

    ratio = str(style.get("ratio", "50/50"))
    parts = ratio.split("/")
    lf = int(parts[0]) / (int(parts[0]) + int(parts[1])) if len(parts) == 2 else 0.5

    usable = SLIDE_WIDTH_EMU - 3 * MARGIN_EMU
    lw = int(usable * lf)
    rw = usable - lw
    rl = MARGIN_EMU + lw + MARGIN_EMU

    # Header content
    hc = slide_data.get("content", {})
    hh = 0
    if hc.get("heading"):
        _add_text_box(pptx_slide, hc["heading"], MARGIN_EMU, MARGIN_EMU,
                      SLIDE_WIDTH_EMU - 2 * MARGIN_EMU, Emu(700000),
                      font_name=ctx.font_heading, font_size=28,
                      bold=True, color_hex=text_color)
        hh = Emu(900000)

    pt = MARGIN_EMU + hh

    # Left panel
    ld = slide_data.get("left", {})
    left_style = ld.get("style", style)
    _add_content_block(pptx_slide, ld.get("content", {}), left_style, ctx,
                       MARGIN_EMU, pt, lw, corpus_root)

    # Right panel
    rd = slide_data.get("right", {})
    right_style = rd.get("style", style)
    _add_content_block(pptx_slide, rd.get("content", {}), right_style, ctx,
                       rl, pt, rw, corpus_root)


def _render_grid(pptx_slide, slide_data: dict, ctx: _ThemeContext, corpus_root: Path,
                 artifact_screenshot: str | None = None):
    from pptx.util import Emu

    style = slide_data.get("style", {})
    cells = slide_data.get("cells", [])
    _set_slide_bg(pptx_slide, style, ctx)

    cols = int(style.get("columns", 2))
    rows = (len(cells) + cols - 1) // max(cols, 1)
    uw = SLIDE_WIDTH_EMU - (cols + 1) * MARGIN_EMU
    uh = SLIDE_HEIGHT_EMU - (max(rows, 1) + 1) * MARGIN_EMU
    cw = uw // max(cols, 1)
    ch = uh // max(rows, 1)

    for idx, cell in enumerate(cells):
        c, r = idx % cols, idx // cols
        x = MARGIN_EMU + c * (cw + MARGIN_EMU)
        y = MARGIN_EMU + r * (ch + MARGIN_EMU)
        _add_content_block(pptx_slide, cell.get("content", {}),
                           cell.get("style", style), ctx, x, y, cw, corpus_root)
