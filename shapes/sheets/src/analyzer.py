"""Deterministic column analysis for CSV/XLSX import. No LLM, no network."""

import csv
import re
from datetime import datetime
from pathlib import Path

BOOLEAN_VALUES = {"true", "false", "1", "0", "yes", "no"}

DATE_FORMATS = [
    ("%Y-%m-%d", "date"),
    ("%m/%d/%Y", "date"),
    ("%d.%m.%Y", "date"),
    ("%Y-%m-%dT%H:%M:%S", "datetime"),
    ("%Y-%m-%dT%H:%M:%SZ", "datetime"),
    ("%m/%d/%Y %H:%M", "datetime"),
    ("%d.%m.%Y %H:%M", "datetime"),
]

URL_PATTERN = re.compile(r"^https?://")

ENUM_MAX_DISTINCT = 30
ENUM_MAX_RATIO = 0.5


def _is_boolean(values: list[str]) -> bool:
    return all(v.lower() in BOOLEAN_VALUES for v in values)


def _is_integer(values: list[str]) -> bool:
    for v in values:
        try:
            int(v)
        except ValueError:
            return False
    return True


def _is_number(values: list[str]) -> bool:
    for v in values:
        try:
            float(v)
        except ValueError:
            return False
    return True


def _try_parse_date(value: str, fmt: str) -> bool:
    try:
        datetime.strptime(value, fmt)
        return True
    except ValueError:
        return False


def _detect_date_format(values: list[str]) -> tuple[str, str] | None:
    """Try each date format. Return (fmt, type_name) if one parses 100% of values."""
    for fmt, type_name in DATE_FORMATS:
        if all(_try_parse_date(v, fmt) for v in values):
            return (fmt, type_name)
    return None


def infer_type(raw_values: list[str | None]) -> tuple[str, str | None]:
    """Infer the narrowest type for a column from all its values.

    Returns (type_name, date_format_or_none).
    100% strictness: any malformed value forces string.
    """
    values = [str(v).strip() for v in raw_values if v is not None and str(v).strip()]
    if not values:
        return ("string", None)

    if _is_boolean(values):
        return ("boolean", None)

    if _is_integer(values):
        return ("integer", None)

    if _is_number(values):
        return ("number", None)

    date_result = _detect_date_format(values)
    if date_result:
        return (date_result[1], date_result[0])

    return ("string", None)


def _count_malformed(values: list[str]) -> int:
    """Count values that would have parsed as a narrower type if the column were clean."""
    best_malformed = len(values)

    for check_fn in [_is_boolean, _is_integer, _is_number]:
        failures = sum(1 for v in values if not check_fn([v]))
        if failures < best_malformed and failures < len(values):
            best_malformed = failures

    for fmt, _ in DATE_FORMATS:
        failures = sum(1 for v in values if not _try_parse_date(v, fmt))
        if failures < best_malformed and failures < len(values):
            best_malformed = failures

    return best_malformed if best_malformed < len(values) else 0


def analyze_column(header: str, position: int, raw_values: list[str | None]) -> dict:
    """Analyze a single column: type, cardinality, enum detection, URL detection."""
    values = [str(v).strip() for v in raw_values if v is not None and str(v).strip()]
    total_count = len(values)
    nulls = len(raw_values) - total_count
    distinct = set(values)
    unique_count = len(distinct)
    is_unique = unique_count == total_count and total_count > 0

    inferred_type, date_format = infer_type(raw_values)

    malformed_count = 0
    if inferred_type == "string" and total_count > 0:
        malformed_count = _count_malformed(values)

    # Enum detection
    enum_values = None
    if (
        inferred_type == "string"
        and unique_count <= ENUM_MAX_DISTINCT
        and total_count > 0
        and unique_count / total_count < ENUM_MAX_RATIO
    ):
        enum_values = sorted(distinct)

    # URL detection: ALL non-null values must be URLs
    looks_like_url = total_count > 0 and all(URL_PATTERN.match(v) for v in values)

    max_length = max((len(v) for v in values), default=0) if values else 0

    samples = values[:5]

    return {
        "header": header,
        "position": position,
        "type": inferred_type,
        "date_format": date_format,
        "nulls": nulls,
        "unique_count": unique_count,
        "total_count": total_count,
        "is_unique": is_unique,
        "values": enum_values,
        "samples": samples,
        "looks_like_url": looks_like_url,
        "max_length": max_length,
        "malformed_count": malformed_count,
    }


def _read_csv(path: Path) -> tuple[list[str], list[list[str]]]:
    """Read CSV file. Returns (headers, rows_as_lists_of_strings)."""
    text = path.read_text(encoding="utf-8-sig")
    reader = csv.reader(text.splitlines())
    rows = list(reader)
    if not rows:
        raise ValueError(f"Empty CSV file: {path}")
    headers = rows[0]
    data_rows = rows[1:]
    return headers, data_rows


def _read_xlsx(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[list[str]]]:
    """Read XLSX file. Returns (headers, rows_as_lists_of_strings)."""
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([str(cell) if cell is not None else "" for cell in row])
    wb.close()
    if not rows:
        raise ValueError(f"Empty XLSX file: {path}")
    headers = rows[0]
    data_rows = rows[1:]
    return headers, data_rows


def analyze_file(
    path: str | Path, sheet_name: str | None = None
) -> dict:
    """Analyze a CSV or XLSX file. Returns the full analysis dict."""
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".csv":
        headers, data_rows = _read_csv(path)
    elif suffix in (".xlsx", ".xls"):
        headers, data_rows = _read_xlsx(path, sheet_name)
    else:
        raise ValueError(f"Unsupported file type: {suffix}. Use .csv or .xlsx")

    columns = []
    for i, header in enumerate(headers):
        col_values = [row[i] if i < len(row) else "" for row in data_rows]
        columns.append(analyze_column(header, i, col_values))

    return {
        "file": str(path),
        "row_count": len(data_rows),
        "columns": columns,
    }
