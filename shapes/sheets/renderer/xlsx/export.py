"""XLSX export renderer.

Receives pre-filtered, pre-ordered Field objects and a resolved Locale.
Reads state.sqlite, writes an XLSX file with formatted columns.
Numbers are written as numeric values with locale-appropriate format strings
so they stay sortable and summable in Excel.
"""
import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from shapes.sheets.src.fields import Field
from shapes.sheets.src.locale import Locale


def export(
    content_dir: Path,
    fields: list[Field],
    export_config: dict,
    output_dir: Path,
    locale: Locale,
) -> list[Path]:
    db_path = content_dir / "state.sqlite"
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    rows = conn.execute("SELECT * FROM rows").fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = content_dir.name
    bold = Font(bold=True)

    # Header row
    for col_idx, field in enumerate(fields, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field.label)
        cell.font = bold

    # Freeze header
    ws.freeze_panes = "A2"

    # Data rows
    for row_idx, row in enumerate(rows, start=2):
        for col_idx, field in enumerate(fields, start=1):
            value = row[field.name]
            cell = ws.cell(row=row_idx, column=col_idx)
            _write_cell(cell, value, field, locale)

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / f"{content_dir.name}.xlsx"
    wb.save(str(out_path))
    return [out_path]


def _write_cell(cell, value, field: Field, locale: Locale) -> None:
    if value is None:
        cell.value = None
        return

    if field.type in ("number", "integer"):
        try:
            cell.value = float(value) if field.type == "number" else int(value)
            cell.number_format = locale.xlsx_number_format
        except (ValueError, TypeError):
            cell.value = str(value)
        return

    if field.type in ("datetime", "date"):
        try:
            dt = datetime.fromisoformat(str(value)).replace(tzinfo=None)
            cell.value = dt
            cell.number_format = locale.xlsx_date_format
        except (ValueError, TypeError):
            cell.value = str(value)
        return

    cell.value = str(value)
